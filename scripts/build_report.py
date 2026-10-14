from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
OUTPUT = ROOT / "7115056075_hw2_report.docx"
FONT = "Microsoft JhengHei"
NAVY = "17365D"
BLUE = "D9EAF7"
LIGHT = "F2F5F7"
GRAY = "666666"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_repeat_table_header_and_style(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index == 0:
                set_cell_shading(cell, NAVY)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, LIGHT)


def set_alt_text(inline_shape, title: str, description: str) -> None:
    drawing = inline_shape._inline
    properties = drawing.docPr
    properties.set("title", title)
    properties.set("descr", description)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    paragraph.add_run(" 頁")


def add_caption(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_figure(document: Document, filename: str, caption: str, description: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = paragraph.add_run().add_picture(str(RESULTS / filename), width=Inches(6.35))
    set_alt_text(shape, caption, description)
    add_caption(document, caption)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.3

    for style_name, size in (("Title", 24), ("Subtitle", 13), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.italic = False
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if style_name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(4)
    title_properties = styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)

    styles["Caption"].font.name = FONT
    styles["Caption"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Caption"].font.color.rgb = RGBColor(80, 80, 80)
    styles["Caption"].font.italic = False

    header = section.header.paragraphs[0]
    header.text = "7115056075｜多元線性回歸與 CRISP-DM"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(section.footer.paragraphs[0])


def add_cover(document: Document, metrics: dict) -> None:
    document.add_paragraph("資料分析實作 HW2", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("零售需求之多元線性回歸分析", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("依 CRISP-DM 完成前 20 名特徵選擇、模型評估與 95% 預測區間", style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("\n")
    meta = document.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    repeat_table_header(meta.rows[0])
    entries = [
        ("學號", "7115056075"),
        ("資料集", "Retail Store Inventory and Demand Forecasting"),
        ("模型", "Multiple Linear Regression"),
        ("完成日期", "2026 年 10 月 14 日"),
    ]
    for row, (key, value) in zip(meta.rows, entries):
        row.cells[0].text = key
        row.cells[1].text = value
        set_cell_shading(row.cells[0], BLUE)
        for cell in row.cells:
            set_cell_margins(cell, top=120, bottom=120)
    document.add_paragraph("\n")
    model = metrics["models"]["Linear regression - top 20"]
    summary = document.add_table(rows=2, cols=4)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    repeat_table_header(summary.rows[0])
    labels = ["選用特徵", "MAE", "RMSE", "R²"]
    values = ["前 20 名", f"{model['MAE']:.2f}", f"{model['RMSE']:.2f}", f"{model['R2']:.4f}"]
    for index, label in enumerate(labels):
        summary.cell(0, index).text = label
        summary.cell(1, index).text = values[index]
        set_cell_shading(summary.cell(0, index), NAVY)
        for run in summary.cell(0, index).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
        summary.cell(1, index).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    note = document.add_paragraph("本報告數值由 7115056075_hw2.ipynb 實際執行產生。")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.runs[0].italic = True
    document.add_page_break()


def add_model_table(document: Document, metrics: dict) -> None:
    table = document.add_table(rows=1, cols=5)
    headers = ["模型", "特徵數", "MAE", "RMSE", "R²"]
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    names = {
        "Mean baseline": "平均值基準",
        "Linear regression - all features": "全部特徵線性回歸",
        "Linear regression - top 20": "前 20 名特徵線性回歸",
    }
    for key, values in metrics["models"].items():
        cells = table.add_row().cells
        row = [names[key], str(values["feature_count"]), f"{values['MAE']:.2f}", f"{values['RMSE']:.2f}", f"{values['R2']:.4f}"]
        for cell, value in zip(cells, row):
            cell.text = value
    set_repeat_table_header_and_style(table)
    add_caption(document, "表 1　測試集模型比較（N = 15,200）")


def add_feature_table(document: Document, features: pd.DataFrame) -> None:
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, ["排名", "特徵", "F-score", "p-value"]):
        cell.text = value
    for _, item in features.iterrows():
        cells = table.add_row().cells
        p_value = "< 0.001" if float(item["P_Value"]) < 0.001 else f"{item['P_Value']:.4f}"
        values = [str(int(item["Rank"])), str(item["Feature"]), f"{item['F_Score']:,.2f}", p_value]
        for cell, value in zip(cells, values):
            cell.text = value
    set_repeat_table_header_and_style(table)
    add_caption(document, "表 2　由訓練集選出的前 20 名特徵")


def build_report(output_path: Path = OUTPUT) -> Path:
    metrics = json.loads((RESULTS / "metrics.json").read_text(encoding="utf-8"))
    features = pd.read_csv(RESULTS / "top20_features.csv")
    document = Document()
    configure_document(document)
    add_cover(document, metrics)

    document.add_heading("摘要", level=1)
    document.add_paragraph(
        "本研究以 Kaggle 的零售庫存與需求合成資料為分析對象，依 CRISP-DM 流程建立多元線性回歸模型。"
        "資料共 76,000 筆、16 欄，扣除日期與目標 Demand 後有 14 個原始候選特徵；類別欄位經 One-Hot Encoding 後形成 44 個模型欄位。"
        "本研究先切分訓練集與測試集，再以訓練集擬合前處理與 F-test 特徵選擇器，選取重要性前 20 名特徵。"
        "前 20 名模型在 15,200 筆測試資料的 MAE 為 16.55、RMSE 為 22.20、R² 為 0.7768；95% 預測區間涵蓋率為 95.60%。"
        "結果顯示縮減特徵後仍保留接近全部 44 個欄位模型的預測能力，但合成資料、隨機切分與可能的目標資訊重疊限制了實務推論。"
    )

    document.add_heading("1. 資料來源與研究脈絡", level=1)
    document.add_paragraph("資料集名稱：Retail Store Inventory and Demand Forecasting")
    document.add_paragraph("來源連結：https://www.kaggle.com/datasets/atomicd/retail-store-inventory-and-demand-forecasting")
    document.add_paragraph(
        "資料涵蓋日期、商店、商品、類別、地區、庫存、銷售量、訂購量、價格、折扣、天氣、促銷、競品價格、季節、疫情與需求。"
        "Kaggle 頁面將其標示為合成資料，因此本研究定位為方法展示，而非真實零售企業績效估計。"
    )
    add_bullets(document, [
        "研究問題：能否用多個零售特徵預測 Demand，並量化單筆預測的不確定性？",
        "目標變數：Demand。",
        "候選原始特徵：14 個；Date 不納入本次橫斷面線性回歸。",
        "主要限制：Units Sold 與 Demand 的定義可能接近，模型可用於預測比較，但係數不可直接視為因果效果。",
    ])

    document.add_heading("2. CRISP-DM 分析流程", level=1)
    document.add_heading("2.1 Business Understanding", level=2)
    document.add_paragraph(
        "需求預測可協助安排補貨與安全庫存。本作業的成功條件是：模型必須優於平均值基準、以可重現方式完成前 20 名特徵選擇，並同時提供點預測、評估指標及 95% 預測區間。"
    )
    document.add_heading("2.2 Data Understanding", level=2)
    add_bullets(document, [
        "資料規模：76,000 筆、16 欄。",
        "訓練集：60,800 筆；測試集：15,200 筆；random_state = 42。",
        "分析先檢查欄位型態、缺失值、重複值與敘述統計。",
        "資料為合成資料，適合比較方法，不應直接外推至真實商店。",
    ])
    document.add_heading("2.3 Data Preparation", level=2)
    add_numbered(document, [
        "先執行 80/20 train-test split，避免測試資料參與前處理與特徵選擇。",
        "數值欄位以 MinMaxScaler 縮放；類別欄位以 OneHotEncoder(drop='first', handle_unknown='ignore') 轉換。",
        "Promotion 與 Epidemic 為二元欄位，直接保留。",
        "前處理器只在訓練集 fit，再分別 transform 訓練集與測試集。",
    ])
    document.add_heading("2.4 Modeling 與前 20 名特徵選擇", level=2)
    document.add_paragraph(
        "本研究使用 SelectKBest 搭配 f_regression，以每個編碼後特徵對 Demand 的單變量 F-score 排名，並依使用者指定保留重要性前 20 名。"
        "選擇器只以訓練資料擬合。模型包含平均值基準、全部 44 個欄位的 LinearRegression，以及前 20 名欄位的 LinearRegression。"
    )
    document.add_heading("2.5 Evaluation", level=2)
    document.add_paragraph(
        "評估指標包括 MAE、RMSE 與 R²。另以 statsmodels OLS 在相同前 20 名欄位上估計 95% prediction interval。"
        "預測區間描述單一新觀測值可能落入的範圍，因此比平均反應的信賴區間更適合備貨風險溝通。"
    )
    document.add_heading("2.6 Deployment", level=2)
    document.add_paragraph(
        "本次交付可執行 Notebook、指標 JSON、比較 CSV、特徵排名 CSV 與圖表。若要部署至真實環境，應固定欄位格式與前處理流程、監測資料漂移，並改用時間切分重新驗證後才可上線。"
    )

    document.add_heading("3. 模型結果", level=1)
    add_model_table(document, metrics)
    document.add_paragraph(
        "全部特徵模型的 R² 為 0.7845；前 20 名模型為 0.7768，差距約 0.0077。前 20 名模型以較少欄位保留大部分預測能力，符合本作業特徵選擇要求。"
        "相較平均值基準，兩個線性回歸模型的 MAE 與 RMSE 均明顯較低。"
    )
    document.add_paragraph(
        f"前 20 名模型的訓練集 OLS R² 為 {metrics['ols']['R2_train']:.4f}，調整後 R² 為 {metrics['ols']['Adjusted_R2_train']:.4f}；"
        f"整體 F 統計量為 {metrics['ols']['F_statistic']:,.2f}。這些統計量說明線性組合具有整體解釋力，但不代表因果關係。"
    )

    document.add_heading("4. 前 20 名特徵", level=1)
    add_feature_table(document, features)
    document.add_paragraph(
        "Units Sold、Units Ordered 與 Epidemic 的 F-score 居前。F-test 是單變量篩選方法，只衡量單一欄位與目標的線性關聯；它不會自動處理特徵間冗餘，也不能證明業務因果。"
    )

    document.add_heading("5. 預測圖與不確定性", level=1)
    add_figure(document, "actual_vs_predicted.png", "圖 1　測試集實際值與預測值", "散點圖比較測試集實際需求與前 20 名線性回歸預測，並含理想對角線。")
    document.add_paragraph("散點整體沿理想線分布，但高需求區域仍有較大誤差，顯示模型無法完全捕捉極端需求。")
    add_figure(document, "prediction_interval.png", "圖 2　測試樣本的 95% 預測區間", "前 120 筆測試樣本的實際值、預測值與百分之九十五預測區間。")
    document.add_paragraph(
        f"測試集 95% 預測區間涵蓋率為 {metrics['prediction_interval']['coverage'] * 100:.2f}%，平均寬度為 {metrics['prediction_interval']['mean_width']:.2f}。"
        "涵蓋率接近名目 95%，但區間偏寬，使用者不應只依點預測決定庫存。"
    )
    add_figure(document, "residual_diagnostics.png", "圖 3　殘差診斷", "左圖為預測值與殘差，右圖為殘差分布直方圖。")
    document.add_paragraph("殘差診斷用來檢查系統性偏差與誤差分布；圖形顯示仍存在較大的正負殘差，因此實務應再比較非線性模型與時間序列驗證。")

    document.add_heading("6. GPT 輔助內容", level=1)
    document.add_paragraph(
        "GPT 協助內容包括：整理作業規格與 CRISP-DM 章節、設計資料洩漏防護流程、建立前 20 名 F-score 特徵選擇、產生多元線性回歸與 statsmodels 預測區間程式、同步 README 與報告，以及執行測試和圖表檢查。"
        "所有模型分數、特徵排名及圖表皆由 Notebook 實際運算，未由 GPT 捏造。完整對話另附「ChatGPT-需求預測與特徵選擇-20261007-1014.pdf」。"
    )

    document.add_heading("7. NotebookLM 摘要", level=1)
    document.add_paragraph(
        "NotebookLM 研究顯示，零售需求預測通常不只比較單一模型，而會依資料型態選擇線性模型、樹模型或時間序列方法。"
        "線性回歸容易解釋各欄位與需求的關聯，可作為清楚的基準；Ridge 可降低高度相關特徵造成的係數不穩定，Lasso 可透過稀疏係數進行內嵌式特徵選擇。"
        "若需求與促銷、天氣及商品類別之間存在非線性互動，可再比較梯度提升樹。具有日期順序的零售資料則應採時間切分或滾動驗證，避免隨機切分把未來資訊混入訓練資料。"
        "模型評估除了 MAE、RMSE 與 R²，也應呈現預測區間，讓管理者同時掌握點預測與缺貨風險。"
        "這些研究結果支持本作業以多元線性回歸建立可解釋基準，並把正則化、非線性模型及時間驗證列為後續改進方向。"
    )
    note = document.add_paragraph("註：本段為使用者提供之 NotebookLM 研究內容的整理；繳交時應一併保留 NotebookLM 來源或操作證明。")
    note.runs[0].italic = True

    document.add_heading("8. 主流或更優方法比較", level=1)
    comparison = document.add_table(rows=1, cols=4)
    for cell, value in zip(comparison.rows[0].cells, ["方法", "核心概念", "可能優點", "本研究中的限制"]):
        cell.text = value
    method_rows = [
        ("Ridge", "L2 正則化", "降低共線性造成的係數不穩定", "需用交叉驗證選擇正則強度"),
        ("Lasso", "L1 正則化", "係數可縮為 0，提供內嵌式選擇", "結果受縮放與超參數影響"),
        ("HistGradientBoosting", "樹的加法集成", "捕捉非線性與特徵交互作用", "解釋性較低，須額外調參"),
        ("ARIMA／SARIMA", "時間相依與季節性", "適合按商店與商品預測未來需求", "需重整為連續時間序列並採時間切分"),
    ]
    for values in method_rows:
        cells = comparison.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    set_repeat_table_header_and_style(comparison)
    add_caption(document, "表 3　其他主流方法之定性比較")
    document.add_paragraph("表中方法未在本作業中實際訓練，因此不宣稱其分數優於線性回歸。合理的後續實驗應使用相同切分與指標進行公平比較。")

    document.add_heading("9. 限制與結論", level=1)
    add_bullets(document, [
        "合成資料不代表真實商店的資料生成過程與市場衝擊。",
        "隨機切分適合本次回歸作業，但零售預測部署應使用時間切分或滾動驗證。",
        "Units Sold 可能與 Demand 高度重疊，會提高預測分數並限制因果解釋。",
        "F-test 僅檢查單變量線性關聯；後續可比較 Ridge、Lasso 與非線性樹模型。",
    ])
    document.add_paragraph(
        "結論：前 20 名特徵線性回歸在可解釋性、特徵數量與預測表現間取得合理平衡，測試集 R² 為 0.7768，95% 預測區間涵蓋率為 95.60%。"
        "本結果足以完成課堂上的 CRISP-DM 與多元線性回歸示範；若用於實務決策，仍需真實資料、時間驗證及模型監控。"
    )

    document.add_page_break()
    document.add_heading("10. 參考資料", level=1)
    references = [
        "Kaggle. Retail Store Inventory and Demand Forecasting. https://www.kaggle.com/datasets/atomicd/retail-store-inventory-and-demand-forecasting",
        "scikit-learn. SelectKBest. https://scikit-learn.org/stable/modules/generated/sklearn.feature_selection.SelectKBest.html",
        "scikit-learn. f_regression. https://scikit-learn.org/stable/modules/generated/sklearn.feature_selection.f_regression.html",
        "scikit-learn. LinearRegression and linear models. https://scikit-learn.org/stable/modules/linear_model.html",
        "scikit-learn. HistGradientBoostingRegressor. https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.HistGradientBoostingRegressor.html",
        "statsmodels. PredictionResults.summary_frame. https://www.statsmodels.org/stable/generated/statsmodels.regression.linear_model.PredictionResults.summary_frame.html",
    ]
    for number, reference in enumerate(references, start=1):
        document.add_paragraph(f"{number}. {reference}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path


if __name__ == "__main__":
    print(build_report())
